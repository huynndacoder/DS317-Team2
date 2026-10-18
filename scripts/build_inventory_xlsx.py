#!/usr/bin/env python3
"""
Build data_inventory.xlsx from dataset_profiles.json.

One tab per dataset (IEEE-CIS, CCFD-2023, ULB). Vietnamese labels.
Every derived figure is an Excel formula over the scanned raw counts, so the
sheet recalculates if a count is corrected by hand.
"""

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
PROFILES = ROOT / "dataset_profiles.json"
OUTPUT = ROOT / "data_inventory.xlsx"

FONT = "Arial"
BLUE = "0000FF"  # hardcoded scan input
BLACK = "000000"  # formula
RED = "FF0000"  # warning
GREY_FILL = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")
YELLOW_FILL = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")

HEADER_ROW = 17
FIRST_DATA_ROW = HEADER_ROW + 1

DICT_HEADERS = [
    "Tên cột",
    "Bảng nguồn",
    "Nhóm",
    "Kiểu dữ liệu",
    "Số giá trị thiếu",
    "Mẫu số",
    "% thiếu",
    "Số giá trị duy nhất",
    "% duy nhất",
    "Min",
    "Max",
    "Trung bình",
    "Độ lệch chuẩn",
    "Top-3 giá trị",
    "Cảnh báo",
    "Mô tả",
]
COL_WIDTHS = [16, 22, 14, 12, 14, 11, 9, 16, 10, 14, 14, 14, 14, 44, 22, 46]

FLAG_VI = {
    "constant": "Hằng số (1 giá trị)",
    "high_missing": "Thiếu >70%",
    "suspected_id": "Nghi là định danh",
}


def ieee_group(col: str):
    exact = {
        "TransactionID": ("Khóa", "Định danh duy nhất của giao dịch"),
        "isFraud": ("Nhãn", "Nhãn gian lận (0 = hợp lệ, 1 = gian lận)"),
        "TransactionDT": (
            "Thời gian",
            "Độ lệch giây so với mốc gốc (không phải datetime)",
        ),
        "TransactionAmt": ("Số tiền", "Số tiền giao dịch (USD) — dùng làm chi phí FN"),
        "ProductCD": ("Sản phẩm", "Mã loại sản phẩm (W/H/C/S/R)"),
        "addr1": (
            "Địa chỉ",
            "Mã vùng địa chỉ thanh toán — thành phần của UID đại diện",
        ),
        "addr2": ("Địa chỉ", "Mã quốc gia địa chỉ thanh toán"),
        "dist1": ("Địa chỉ", "Khoảng cách (billing vs địa chỉ khác)"),
        "dist2": ("Địa chỉ", "Khoảng cách (biến thể thứ hai)"),
        "P_emaildomain": ("Email", "Tên miền email người mua — nút đồ thị tiềm năng"),
        "R_emaildomain": ("Email", "Tên miền email người nhận — nút đồ thị tiềm năng"),
        "DeviceType": ("Thiết bị", "Loại thiết bị (desktop/mobile)"),
        "DeviceInfo": ("Thiết bị", "Chuỗi mô tả thiết bị — nút đồ thị tiềm năng"),
        "card1": ("Thẻ", "Định danh thẻ — thành phần chính của UID đại diện"),
        "card2": ("Thẻ", "Thuộc tính thẻ (nhiều khả năng là BIN)"),
        "card3": ("Thẻ", "Thuộc tính thẻ"),
        "card4": ("Thẻ", "Tổ chức phát hành (visa/mastercard/...)"),
        "card5": ("Thẻ", "Thuộc tính thẻ"),
        "card6": ("Thẻ", "Loại thẻ (credit/debit)"),
        "id_30": ("Identity", "Hệ điều hành — nút đồ thị tiềm năng"),
        "id_31": ("Identity", "Trình duyệt — nút đồ thị tiềm năng"),
    }
    if col in exact:
        return exact[col]
    if col.startswith("card"):
        return ("Thẻ", "Định danh/thuộc tính thẻ")
    if col.startswith("C"):
        return (
            "C-đếm",
            "Biến đếm của Vesta (đếm địa chỉ/thẻ liên kết) — mang bản chất đồ thị",
        )
    if col.startswith("D"):
        return ("D-thời gian", "Số ngày kể từ một sự kiện trước đó")
    if col.startswith("M"):
        return ("M-cờ khớp", "Cờ khớp thông tin (tên trên thẻ vs địa chỉ…)")
    if col.startswith("V"):
        return (
            "V-Vesta",
            "Đặc trưng kỹ nghệ của Vesta (đếm, xếp hạng, quan hệ) — ý nghĩa không công bố",
        )
    if col.startswith("id_"):
        return ("Identity", "Thông tin định danh/phiên (ý nghĩa không công bố)")
    return ("Khác", "")


def _pca_group(col: str, time_desc: str):
    if col in ("id", "Time"):
        return ("Khóa/Thời gian", time_desc)
    if col == "Amount":
        return ("Số tiền", "Số tiền giao dịch")
    if col == "Class":
        return ("Nhãn", "Nhãn gian lận (0 = hợp lệ, 1 = gian lận)")
    if col.startswith("V"):
        return ("PCA", "Thành phần chính sau PCA — KHÔNG thể suy ngược ra thực thể")
    return ("Khác", "")


def ccfd_group(col: str):
    return _pca_group(col, "Chỉ số dòng, không mang thông tin thời gian")


def ulb_group(col: str):
    return _pca_group(col, "Độ lệch giây kể từ giao dịch đầu tiên (trải 2 ngày)")


def _put(
    ws,
    row,
    col,
    value,
    *,
    bold=False,
    color=BLACK,
    size=10,
    fmt=None,
    fill=None,
    wrap=False,
):
    cell = ws.cell(row, col, value)
    cell.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def build_sheet(ws, prof, meta):
    n_rows = prof["n_rows"]

    # ---- header block -------------------------------------------------
    rows = [
        ("Tên bộ dữ liệu", prof["display_name"], BLUE, None),
        ("Nguồn", meta["url"], BLUE, None),
        ("Tệp đã rà soát", ", ".join(prof["files"]), BLUE, None),
        ("Vai trò trong đồ án", meta["role"], BLUE, None),
        ("Số dòng", n_rows, BLUE, "#,##0"),
        ("Số cột", meta["cols_text"](prof), BLUE, None),
        ("Dung lượng trên đĩa (MB)", prof["bytes_on_disk"] / 1e6, BLUE, "#,##0.0"),
        ("Cột nhãn", prof["label_col"], BLUE, None),
        ("Số ca gian lận", prof["n_positive"], BLUE, "#,##0"),
    ]
    for i, (label, value, color, fmt) in enumerate(rows, start=1):
        _put(ws, i, 1, label, bold=True)
        _put(ws, i, 2, value, color=color, fmt=fmt)

    # Fraud rate is derived, so it is a formula over B5/B9 rather than a
    # number typed in from the paper.
    _put(ws, 10, 1, "Tỷ lệ gian lận", bold=True)
    _put(ws, 10, 2, "=IFERROR(B9/B5,0)", fmt="0.000%")

    tail = [
        ("Ẩn danh hóa", meta["anon"]),
        ("Giữ định danh thực thể?", meta["entities"]),
        ("Có trục thời gian?", meta["timeline"]),
        ("Dựng được đồ thị giao dịch?", meta["graph"]),
        ("Ngày rà soát", date.today().isoformat()),
    ]
    for i, (label, value) in enumerate(tail, start=11):
        _put(ws, i, 1, label, bold=True)
        _put(ws, i, 2, value, color=BLUE)

    _put(ws, 16, 1, "Ghi chú", bold=True)
    _put(ws, 16, 2, meta["note"], fill=YELLOW_FILL, wrap=True)
    ws.row_dimensions[16].height = 30

    # ---- dictionary ---------------------------------------------------
    for c, header in enumerate(DICT_HEADERS, start=1):
        _put(ws, HEADER_ROW, c, header, bold=True, fill=GREY_FILL, wrap=True)
    ws.row_dimensions[HEADER_ROW].height = 30

    row = FIRST_DATA_ROW
    for name, info in prof["columns"].items():
        group, desc = meta["group_fn"](name)

        _put(ws, row, 1, name, color=BLUE)
        _put(ws, row, 2, info["source_table"], color=BLUE, size=9)
        _put(ws, row, 3, group, color=BLUE)
        _put(ws, row, 4, info["dtype"], color=BLUE)
        _put(ws, row, 5, info["n_missing"], color=BLUE, fmt="#,##0")
        # Identity columns are measured against identity rows, not the
        # 590,540 transaction rows, so the denominator is explicit per row.
        _put(ws, row, 6, info["denominator"], color=BLUE, fmt="#,##0")
        _put(ws, row, 7, f"=IFERROR(E{row}/F{row},0)", fmt="0.0%")
        _put(ws, row, 8, info["n_unique"], color=BLUE, fmt="#,##0")
        _put(ws, row, 9, f"=IFERROR(H{row}/F{row},0)", fmt="0.0%")

        for col_idx, key in ((10, "min"), (11, "max"), (12, "mean"), (13, "std")):
            value = info.get(key)
            if value is not None:
                _put(ws, row, col_idx, value, color=BLUE, fmt="#,##0.0000")

        top3 = info.get("top3") or []
        if top3:
            text = "; ".join(f"{v} ({c:,})" for v, c in top3)
            _put(ws, row, 14, text, color=BLUE, size=9)

        flags = info.get("flags") or []
        if flags:
            _put(
                ws,
                row,
                15,
                ", ".join(FLAG_VI.get(f, f) for f in flags),
                color=RED,
                size=9,
            )

        _put(ws, row, 16, desc, size=9, wrap=True)
        row += 1

    last = row - 1

    # ---- summary block, formulas over the dictionary -------------------
    summary_col = len(DICT_HEADERS) + 2  # leave one blank column
    letter_flag = get_column_letter(15)
    letter_pct = get_column_letter(7)
    rng_flag = f"{letter_flag}{FIRST_DATA_ROW}:{letter_flag}{last}"
    rng_pct = f"{letter_pct}{FIRST_DATA_ROW}:{letter_pct}{last}"

    _put(ws, 1, summary_col, "Tổng hợp", bold=True)
    summary = [
        ("Tổng số cột", f"=COUNTA(A{FIRST_DATA_ROW}:A{last})", "#,##0"),
        ("Số cột thiếu >70%", f'=COUNTIF({rng_flag},"*Thiếu >70%*")', "#,##0"),
        ("Số cột hằng số", f'=COUNTIF({rng_flag},"*Hằng số*")', "#,##0"),
        (
            "Số cột nghi định danh",
            f'=COUNTIF({rng_flag},"*Nghi là định danh*")',
            "#,##0",
        ),
        ("Số cột không thiếu giá trị", f"=COUNTIF({rng_pct},0)", "#,##0"),
        ("% thiếu trung bình", f"=IFERROR(AVERAGE({rng_pct}),0)", "0.0%"),
        ("% thiếu lớn nhất", f"=IFERROR(MAX({rng_pct}),0)", "0.0%"),
    ]
    for i, (label, formula, fmt) in enumerate(summary, start=2):
        _put(ws, i, summary_col, label, bold=True)
        _put(ws, i, summary_col + 1, formula, fmt=fmt)

    _put(ws, 10, summary_col, "Chú giải màu", bold=True)
    for i, (text, color) in enumerate(
        [
            ("Xanh dương: số liệu quét trực tiếp từ CSV", BLUE),
            ("Đen: công thức Excel tự tính", BLACK),
            ("Đỏ: cảnh báo chất lượng", RED),
        ],
        start=11,
    ):
        _put(ws, i, summary_col, text, color=color, size=9)

    # ---- layout --------------------------------------------------------
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(DICT_HEADERS))}{last}"
    for c, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.column_dimensions[get_column_letter(summary_col)].width = 28
    ws.column_dimensions[get_column_letter(summary_col + 1)].width = 12
    return last


def main():
    profiles = json.loads(PROFILES.read_text())

    metas = {
        "ieee_cis": dict(
            tab="IEEE-CIS",
            url="https://www.kaggle.com/competitions/ieee-fraud-detection/data",
            role="CHÍNH — dùng cho toàn bộ thực nghiệm",
            anon="Không — giữ nguyên trường nghiệp vụ",
            entities="CÓ: card1–card6, addr1/addr2, P/R_emaildomain, DeviceInfo, id_30, id_31",
            timeline="CÓ: TransactionDT (độ lệch giây, trải ~6 tháng)",
            graph="CÓ — đây là lý do bộ này được chọn làm bộ chính",
            note=(
                "Chỉ rà soát tập train. Cột identity chỉ phủ 144.233/590.540 giao dịch "
                "(24,4%), nên %thiếu của nhóm id_*/Device* tính trên mẫu số 144.233. "
                "IEEE-CIS KHÔNG có cột UID — UID đại diện (card1+addr1+D1n) là GIẢ ĐỊNH của nhóm."
            ),
            cols_text=lambda p: (
                f"{p['n_cols_txn']} (transaction) + {p['n_cols_ident']} "
                f"(identity) = {len(p['columns'])} cột phân biệt"
            ),
            group_fn=ieee_group,
        ),
        "ccfd_2023": dict(
            tab="CCFD-2023",
            url="https://www.kaggle.com/datasets/nelgiriyewithana/credit-card-fraud-detection-dataset-2023",
            role="PHỤ — nhóm đối chứng để bác bỏ accuracy",
            anon="ĐÃ PCA: V1–V28",
            entities="KHÔNG — mọi định danh đã bị PCA xóa",
            timeline="KHÔNG — cột id chỉ là chỉ số dòng",
            graph="KHÔNG — không còn thực thể để làm nút",
            note=(
                "Tỷ lệ 50/50 KHÔNG phải tỷ lệ gian lận thực tế mà do lấy mẫu lại. "
                "Không dùng bộ này để công bố bất kỳ con số hiệu năng chính nào."
            ),
            cols_text=lambda p: f"{p['n_cols']} cột",
            group_fn=ccfd_group,
        ),
        "ulb": dict(
            tab="ULB",
            url="https://www.kaggle.com/datasets/mlg-ulb/creditcardfraud",
            role="PHỤ — kiểm thử mất cân bằng cực đoan",
            anon="ĐÃ PCA: V1–V28",
            entities="KHÔNG — mọi định danh đã bị PCA xóa",
            timeline="CÓ: Time (độ lệch giây, trải 2 ngày)",
            graph="KHÔNG — không còn thực thể để làm nút",
            note=(
                "492/284.807 ca gian lận (1/579). Dùng để kiểm chứng cost-sensitive "
                "learning trong kịch bản của P1 (Chan et al., 1999)."
            ),
            cols_text=lambda p: f"{p['n_cols']} cột",
            group_fn=ulb_group,
        ),
    }

    wb = Workbook()
    wb.remove(wb.worksheets[0])

    for key, meta in metas.items():
        ws = wb.create_sheet(meta["tab"])
        last = build_sheet(ws, profiles[key], meta)
        print(f"  {meta['tab']}: {last - FIRST_DATA_ROW + 1} cột")

    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
